import io
import unittest
from types import SimpleNamespace

from docx import Document
from reportlab.pdfgen import canvas

from rag import (
    Chunk,
    IndexedChunk,
    TextBlock,
    TokenUsage,
    answer_question,
    index_chunks,
    load_document,
    make_chunks,
    retrieve,
)


class DocumentTests(unittest.TestCase):
    def test_pdf_extracts_text_with_page_number(self):
        data = io.BytesIO()
        pdf = canvas.Canvas(data)
        pdf.drawString(72, 700, "Python backend engineer")
        pdf.showPage()
        pdf.drawString(72, 700, "SQL Server reporting")
        pdf.save()

        blocks = load_document("cv.pdf", data.getvalue())

        self.assertEqual(
            [(block.source, block.text.strip()) for block in blocks],
            [
                ("PDF page 1", "Python backend engineer"),
                ("PDF page 2", "SQL Server reporting"),
            ],
        )

    def test_docx_extracts_paragraphs_and_table_cells(self):
        document = Document()
        document.add_paragraph("Backend engineer")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Python"
        table.cell(0, 1).text = "SQL Server"
        data = io.BytesIO()
        document.save(data)

        blocks = load_document("cv.docx", data.getvalue())

        self.assertEqual(
            [block.text for block in blocks],
            ["Backend engineer", "Python", "SQL Server"],
        )
        self.assertTrue(all(block.source for block in blocks))

    def test_empty_document_is_rejected_before_embedding(self):
        document = Document()
        data = io.BytesIO()
        document.save(data)

        with self.assertRaisesRegex(ValueError, "extractable text"):
            load_document("empty.docx", data.getvalue())


class RetrievalTests(unittest.TestCase):
    def test_chunks_preserve_text_and_source_locations(self):
        blocks = [
            TextBlock("Python API integrations and services", "PDF page 1"),
            TextBlock("SQL Server reporting and analysis", "PDF page 2"),
        ]

        chunks = make_chunks(blocks, max_chars=45, overlap=8)

        self.assertGreaterEqual(len(chunks), 2)
        self.assertTrue(any("Python" in chunk.text for chunk in chunks))
        self.assertTrue(any("SQL Server" in chunk.text for chunk in chunks))
        self.assertTrue(any("PDF page 1" in chunk.source for chunk in chunks))
        self.assertTrue(any("PDF page 2" in chunk.source for chunk in chunks))

    def test_retrieve_returns_nearest_chunk_first(self):
        index = [
            IndexedChunk(Chunk("1", "Python experience", "PDF page 1"), (1.0, 0.0)),
            IndexedChunk(Chunk("2", "SQL experience", "PDF page 2"), (0.0, 1.0)),
        ]

        hits = retrieve(index, (0.9, 0.1), limit=2)

        self.assertEqual([hit.chunk.id for hit in hits], ["1", "2"])
        self.assertGreater(hits[0].score, hits[1].score)


class FakeOpenAI:
    def __init__(self):
        self.embedding_calls = []
        self.response_calls = []
        self.embeddings = SimpleNamespace(create=self.create_embedding)
        self.responses = SimpleNamespace(create=self.create_response)

    def create_embedding(self, **kwargs):
        self.embedding_calls.append(kwargs)
        inputs = kwargs["input"]
        if isinstance(inputs, str):
            inputs = [inputs]
        data = []
        for index, value in enumerate(inputs):
            vector = [1.0, 0.0] if "Python" in value else [0.0, 1.0]
            data.append(SimpleNamespace(index=index, embedding=vector))
        return SimpleNamespace(data=data, usage=SimpleNamespace(prompt_tokens=17))

    def create_response(self, **kwargs):
        self.response_calls.append(kwargs)
        return SimpleNamespace(
            output_text="มีประสบการณ์ Python [1]",
            usage=SimpleNamespace(
                input_tokens=100,
                output_tokens=20,
                input_tokens_details=SimpleNamespace(cached_tokens=5, cache_write_tokens=10),
            ),
        )


class OpenAITests(unittest.TestCase):
    def test_usage_counts_embedding_and_answer_tokens(self):
        client = FakeOpenAI()
        usage = TokenUsage()

        index = index_chunks(
            client, [Chunk("1", "Python APIs", "PDF page 1")], usage=usage
        )
        answer_question(client, index, "What Python experience?", usage=usage)

        self.assertEqual(usage.embedding_tokens, 34)
        self.assertEqual(usage.llm_input_tokens, 100)
        self.assertEqual(usage.llm_cached_input_tokens, 5)
        self.assertEqual(usage.llm_cache_write_tokens, 10)
        self.assertEqual(usage.llm_output_tokens, 20)

    def test_estimated_cost_uses_standard_rates_and_cache_categories(self):
        usage = TokenUsage(
            embedding_tokens=1000,
            llm_input_tokens=2000,
            llm_cached_input_tokens=100,
            llm_cache_write_tokens=50,
            llm_output_tokens=1000,
        )

        self.assertAlmostEqual(usage.estimated_usd, 0.00071225)

    def test_email_question_recalls_email_without_unrelated_chunks(self):
        class EmailQuestionClient(FakeOpenAI):
            def create_embedding(self, **kwargs):
                self.embedding_calls.append(kwargs)
                return SimpleNamespace(data=[SimpleNamespace(index=0, embedding=[0.0, 1.0])])

        client = EmailQuestionClient()
        index = [
            IndexedChunk(Chunk("1", "jane@example.com", "PDF page 1"), (1.0, 0.0)),
            IndexedChunk(Chunk("2", "Backend systems", "PDF page 1"), (0.0, 1.0)),
            IndexedChunk(Chunk("3", "SQL projects", "PDF page 2"), (0.0, 1.0)),
            IndexedChunk(Chunk("4", "Python projects", "PDF page 2"), (0.0, 1.0)),
        ]

        _, hits = answer_question(client, index, "อีเมลเจ้าของ CV คืออะไร")

        self.assertEqual([hit.chunk.id for hit in hits], ["1"])
        self.assertIn("jane@example.com", client.response_calls[0]["input"])
        self.assertNotIn("Backend systems", client.response_calls[0]["input"])

    def test_contact_question_includes_phone_and_email_from_separate_chunks(self):
        class ContactQuestionClient(FakeOpenAI):
            def create_embedding(self, **kwargs):
                self.embedding_calls.append(kwargs)
                return SimpleNamespace(data=[SimpleNamespace(index=0, embedding=[0.0, 1.0])])

        client = ContactQuestionClient()
        index = [
            IndexedChunk(Chunk("1", "Phone: 081-234-5678", "PDF page 1"), (1.0, 0.0)),
            IndexedChunk(Chunk("2", "jane@example.com", "PDF page 1"), (1.0, 0.0)),
            IndexedChunk(Chunk("3", "Backend systems", "PDF page 2"), (0.0, 1.0)),
            IndexedChunk(Chunk("4", "SQL projects", "PDF page 2"), (0.0, 1.0)),
            IndexedChunk(Chunk("5", "Python projects", "PDF page 2"), (0.0, 1.0)),
        ]

        _, hits = answer_question(client, index, "เบอร์โทรและอีเมลคืออะไร")

        self.assertEqual({hit.chunk.id for hit in hits}, {"1", "2"})
        self.assertIn("081-234-5678", client.response_calls[0]["input"])
        self.assertIn("jane@example.com", client.response_calls[0]["input"])

    def test_exact_skill_question_recalls_chunk_missed_by_embedding(self):
        class SkillQuestionClient(FakeOpenAI):
            def create_embedding(self, **kwargs):
                self.embedding_calls.append(kwargs)
                return SimpleNamespace(data=[SimpleNamespace(index=0, embedding=[0.0, 1.0])])

        client = SkillQuestionClient()
        index = [
            IndexedChunk(Chunk("1", "Built Python APIs", "PDF page 1"), (1.0, 0.0)),
            IndexedChunk(Chunk("2", "Managed projects", "PDF page 1"), (0.0, 1.0)),
            IndexedChunk(Chunk("3", "Led a team", "PDF page 2"), (0.0, 1.0)),
            IndexedChunk(Chunk("4", "Created reports", "PDF page 2"), (0.0, 1.0)),
        ]

        _, hits = answer_question(client, index, "มีประสบการณ์ Python อะไรบ้าง")

        self.assertEqual(hits[0].chunk.id, "1")
        self.assertLessEqual(len(hits), 3)
        self.assertEqual(len(client.response_calls), 1)

    def test_phone_question_does_not_treat_unlabeled_id_as_phone(self):
        class PhoneQuestionClient(FakeOpenAI):
            def create_embedding(self, **kwargs):
                return SimpleNamespace(data=[SimpleNamespace(index=0, embedding=[0.0, 1.0])])

        index = [
            IndexedChunk(Chunk("1", "ID 1234567890123", "PDF page 1"), (1.0, 0.0)),
            IndexedChunk(Chunk("2", "Backend systems", "PDF page 1"), (0.0, 1.0)),
            IndexedChunk(Chunk("3", "SQL projects", "PDF page 2"), (0.0, 1.0)),
            IndexedChunk(Chunk("4", "Python projects", "PDF page 2"), (0.0, 1.0)),
        ]

        _, hits = answer_question(PhoneQuestionClient(), index, "เจ้าของ CV นี้เบอร์โทรอะไร")

        self.assertEqual([hit.chunk.id for hit in hits], ["2", "3", "4"])

    def test_phone_question_includes_contact_chunk_when_embedding_misses_it(self):
        class PhoneQuestionClient(FakeOpenAI):
            def create_embedding(self, **kwargs):
                self.embedding_calls.append(kwargs)
                return SimpleNamespace(
                    data=[SimpleNamespace(index=0, embedding=[0.0, 1.0])]
                )

            def create_response(self, **kwargs):
                self.response_calls.append(kwargs)
                answer = (
                    "เบอร์โทรคือ 081-234-5678 [1]"
                    if "081-234-5678" in kwargs["input"]
                    else "ไม่พบข้อมูลนี้ใน CV"
                )
                return SimpleNamespace(output_text=answer)

        client = PhoneQuestionClient()
        index = [
            IndexedChunk(Chunk("1", "Phone: 081-234-5678", "PDF page 1"), (1.0, 0.0)),
            IndexedChunk(Chunk("2", "Backend systems", "PDF page 1"), (0.0, 1.0)),
            IndexedChunk(Chunk("3", "SQL projects", "PDF page 2"), (0.0, 1.0)),
            IndexedChunk(Chunk("4", "Python projects", "PDF page 2"), (0.0, 1.0)),
        ]

        answer, hits = answer_question(client, index, "เจ้าของ CV นี้เบอร์โทรอะไร")

        self.assertIn("081-234-5678", answer)
        self.assertEqual([hit.chunk.id for hit in hits], ["1"])
        self.assertIn("081-234-5678", client.response_calls[0]["input"])
        self.assertNotIn("Backend systems", client.response_calls[0]["input"])

    def test_index_and_answer_use_requested_models_and_cite_retrieved_cv(self):
        client = FakeOpenAI()
        chunks = [
            Chunk("1", "Python API experience", "PDF page 1"),
            Chunk("2", "SQL reporting", "PDF page 2"),
        ]

        index = index_chunks(client, chunks)
        answer, hits = answer_question(client, index, "What Python experience?")

        self.assertEqual(client.embedding_calls[0]["model"], "text-embedding-3-small")
        self.assertEqual(client.embedding_calls[0]["input"], [c.text for c in chunks])
        self.assertEqual(client.embedding_calls[1]["model"], "text-embedding-3-small")
        self.assertEqual([hit.chunk.id for hit in hits][0], "1")
        self.assertIn("Python", answer)
        self.assertEqual(client.response_calls[0]["model"], "gpt-6-luna")
        self.assertIs(client.response_calls[0]["store"], False)
        self.assertIn("PDF page 1", client.response_calls[0]["input"])
        self.assertIn("ไม่พบข้อมูล", client.response_calls[0]["instructions"])

    def test_no_index_does_not_call_openai(self):
        client = FakeOpenAI()

        answer, hits = answer_question(client, [], "What is my GPA?")

        self.assertEqual(answer, "ไม่พบข้อมูลนี้ใน CV")
        self.assertEqual(hits, [])
        self.assertEqual(client.embedding_calls, [])
        self.assertEqual(client.response_calls, [])


if __name__ == "__main__":
    unittest.main()
